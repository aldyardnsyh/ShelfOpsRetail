"""Document parsing for MD/TXT (native), PDF (pypdf), DOCX (python-docx).

The parser returns heading-aware blocks so the chunker never merges two
sections into one chunk. Page numbers are preserved for citations.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_MIMES = frozenset(
    {
        "text/markdown",
        "text/x-markdown",
        "text/plain",
        "text/md",
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
)

_MIME_BY_SUFFIX = {
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_HEADING_INLINE_RE = re.compile(r"\s+#+\s*$")


class ParseError(ValueError):
    """Raised when a document cannot be parsed."""


@dataclass(frozen=True)
class TextBlock:
    text: str
    heading_path: str = ""
    page: int | None = None


@dataclass(frozen=True)
class ParsedDocument:
    blocks: list[TextBlock] = field(default_factory=list)
    mime: str = "text/plain"
    page_count: int = 0

    @property
    def text(self) -> str:
        return "\n\n".join(block.text for block in self.blocks)


def detect_mime(mime: str, filename: str | None = None) -> str:
    if mime and mime in SUPPORTED_MIMES:
        return "text/markdown" if mime == "text/x-markdown" else mime
    if filename:
        suffix = Path(filename).suffix.lower()
        if suffix in _MIME_BY_SUFFIX:
            return _MIME_BY_SUFFIX[suffix]
    raise ParseError(f"unsupported document type: mime={mime!r} filename={filename!r}")


def parse_document(
    data: bytes,
    mime: str,
    *,
    filename: str | None = None,
) -> ParsedDocument:
    if not data:
        raise ParseError("document is empty")
    resolved = detect_mime(mime, filename)
    if resolved == "text/markdown":
        return _parse_markdown(_decode(data), resolved)
    if resolved == "text/plain":
        return _parse_text(_decode(data), resolved)
    if resolved == "application/pdf":
        return _parse_pdf(data, resolved)
    if resolved.endswith("wordprocessingml.document"):
        return _parse_docx(data, resolved)
    raise ParseError(f"unsupported mime: {resolved}")


def _decode(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")


def _parse_text(text: str, mime: str) -> ParsedDocument:
    cleaned = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not cleaned:
        raise ParseError("document contains no text")
    return ParsedDocument(blocks=[TextBlock(text=cleaned)], mime=mime, page_count=1)


def _parse_markdown(text: str, mime: str) -> ParsedDocument:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[TextBlock] = []
    heading_stack: list[str] = []
    buffer: list[str] = []
    in_fence = False

    def flush() -> None:
        body = "\n".join(buffer).strip()
        buffer.clear()
        if body:
            blocks.append(TextBlock(text=body, heading_path=" > ".join(heading_stack)))

    for line in lines:
        stripped = line.strip()
        if stripped.startswith(("```", "~~~")):
            in_fence = not in_fence
            buffer.append(line)
            continue
        if not in_fence:
            match = _HEADING_RE.match(stripped)
            if match:
                flush()
                level = len(match.group(1))
                title = _HEADING_INLINE_RE.sub("", match.group(2)).strip()
                heading_stack = heading_stack[: level - 1]
                while len(heading_stack) < level - 1:
                    heading_stack.append("")
                heading_stack.append(title)
                continue
        buffer.append(line)

    flush()
    if not blocks:
        raise ParseError("markdown contains no body text")
    return ParsedDocument(blocks=blocks, mime=mime, page_count=1)


def _parse_pdf(data: bytes, mime: str) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        blocks: list[TextBlock] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                blocks.append(TextBlock(text=text, page=index))
        return ParsedDocument(blocks=blocks, mime=mime, page_count=len(reader.pages))
    except ParseError:
        raise
    except Exception as exc:  # pypdf raises many internal error types
        raise ParseError(f"failed to parse PDF: {exc}") from exc


def _parse_docx(data: bytes, mime: str) -> ParsedDocument:
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"failed to parse DOCX: {exc}") from exc

    blocks: list[TextBlock] = []
    heading_stack: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "") or ""
        if style.lower().startswith("heading"):
            level = _docx_heading_level(style)
            heading_stack = heading_stack[: level - 1]
            while len(heading_stack) < level - 1:
                heading_stack.append("")
            heading_stack.append(text)
            continue
        blocks.append(TextBlock(text=text, heading_path=" > ".join(heading_stack)))

    if not blocks:
        raise ParseError("DOCX contains no paragraph text")
    return ParsedDocument(blocks=blocks, mime=mime, page_count=0)


def _docx_heading_level(style: str) -> int:
    match = re.search(r"(\d+)", style)
    if not match:
        return 1
    return max(1, min(6, int(match.group(1))))
