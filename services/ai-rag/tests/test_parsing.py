from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument

from app.parsing import ParseError, detect_mime, parse_document


def test_markdown_is_heading_aware():
    text = (
        "Pengantar tanpa heading.\n\n"
        "# Operasional\n"
        "Stok minimum lima unit.\n\n"
        "## Harga\n"
        "Harga promo berlaku tujuh hari.\n"
    )
    parsed = parse_document(text.encode("utf-8"), "text/markdown")

    assert [block.heading_path for block in parsed.blocks] == [
        "",
        "Operasional",
        "Operasional > Harga",
    ]
    assert parsed.blocks[1].text == "Stok minimum lima unit."


def test_markdown_ignores_headings_inside_code_fence():
    text = "# Judul\n```\n# bukan heading\n```\nIsi nyata.\n"
    parsed = parse_document(text.encode("utf-8"), "text/markdown")
    assert len(parsed.blocks) == 1
    assert parsed.blocks[0].heading_path == "Judul"
    assert "# bukan heading" in parsed.blocks[0].text


def test_plain_text():
    parsed = parse_document(b"Stok minimum lima unit.\r\nHarga stabil.", "text/plain")
    assert parsed.page_count == 1
    assert parsed.blocks[0].text == "Stok minimum lima unit.\nHarga stabil."


def test_docx_headings_and_paragraphs():
    document = DocxDocument()
    document.add_heading("Operasional", level=1)
    document.add_paragraph("Stok minimum lima unit.")
    document.add_heading("Harga", level=2)
    document.add_paragraph("Harga promo tujuh hari.")
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse_document(
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert [block.heading_path for block in parsed.blocks] == [
        "Operasional",
        "Operasional > Harga",
    ]
    assert "lima unit" in parsed.blocks[0].text


def test_pdf_page_text_and_number():
    parsed = parse_document(_minimal_pdf("Stok minimum lima unit."), "application/pdf")
    assert parsed.page_count == 1
    assert parsed.blocks[0].page == 1
    assert "Stok minimum lima unit." in parsed.blocks[0].text


def test_detect_mime_from_filename():
    assert detect_mime("", "kebijakan.md") == "text/markdown"
    assert detect_mime("", "data.txt") == "text/plain"
    assert detect_mime("application/pdf", "x.bin") == "application/pdf"


def test_unsupported_type_rejected():
    with pytest.raises(ParseError):
        parse_document(b"abc", "image/png", filename="x.png")


def test_empty_document_rejected():
    with pytest.raises(ParseError):
        parse_document(b"", "text/plain")


def _minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)
